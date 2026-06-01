"""Export a readable Vietnamese narrative from a Flow V2 rebalance journal."""

from __future__ import annotations

import argparse
from html import escape
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def _date_vn(value: object) -> str:
    return pd.to_datetime(value).strftime("%d/%m/%Y")


def _money_vn(value: float) -> str:
    amount = abs(float(value))
    if amount >= 1_000_000_000:
        shown = f"{amount / 1_000_000_000:.3f} ty dong"
    elif amount >= 1_000_000:
        shown = f"{amount / 1_000_000:.2f} trieu dong"
    else:
        shown = f"{amount:,.0f} dong"
    return shown


def _result_vn(value: float) -> str:
    if value > 0:
        return f"lai {_money_vn(value)}"
    if value < 0:
        return f"lo {_money_vn(value)}"
    return "khong phat sinh lai lo"


def _with_accents(value: str) -> str:
    replacements = {
        "ty dong": "t\u1ef7 \u0111\u1ed3ng",
        "trieu dong": "tri\u1ec7u \u0111\u1ed3ng",
        " dong": " \u0111\u1ed3ng",
        "lai ": "l\u00e3i ",
        "lo ": "l\u1ed7 ",
        "khong phat sinh lai lo": "kh\u00f4ng ph\u00e1t sinh l\u00e3i/l\u1ed7",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    return value


def _max_drawdown_pct(values: pd.Series) -> float:
    values = values.astype(float)
    return float((values / values.cummax() - 1.0).min() * 100.0) if not values.empty else 0.0


def _describe(row: pd.Series, realized_by_date: dict[str, float]) -> str:
    signal_date = _date_vn(row["Ngay_tin_hieu_T"])
    execute_date = _date_vn(row["Ngay_mua_ban_rebalance_T1"])
    end_date = _date_vn(row["Ngay_chot_PnL_ky"])
    pnl = _with_accents(_result_vn(float(row["Lai_lo_rong_ky"])))
    pct = float(row["Ty_suat_ky_pct"])
    percent = f"{pct:+.2f}%" if pct else "0.00%"
    fees = _with_accents(_money_vn(float(row["Phi_thue_ky"])))
    sold_result = realized_by_date.get(str(row["Ngay_mua_ban_rebalance_T1"]))
    realized = ""
    if sold_result is not None:
        realized = f"; ri\u00eang ph\u1ea7n b\u00e1n ra \u0111\u00e3 ch\u1ed1t {_with_accents(_result_vn(sold_result))}"
    if row["Loai_quyet_dinh"] == "BAN_VE_TIEN_MAT":
        action = (
            f"gate kh\u00f4ng c\u00f2n cho ph\u00e9p n\u1eafm gi\u1eef c\u1ed5 phi\u1ebfu; ng\u00e0y {execute_date}, "
            f"h\u1ec7 th\u1ed1ng b\u00e1n {row['Ma_ban_ra']} v\u00e0 chuy\u1ec3n danh m\u1ee5c v\u1ec1 100% ti\u1ec1n m\u1eb7t, "
            f"gi\u00e1 tr\u1ecb b\u00e1n {_with_accents(_money_vn(float(row['Gia_tri_ban'])))}; chi ph\u00ed {fees}{realized}"
        )
    elif row["Ma_ban_ra"] != "-" and row["Ma_mua_them"] != "-":
        target_symbols = {item.strip() for item in str(row["Danh_muc_muc_tieu"]).split(",")}
        sold_symbols = {item.strip() for item in str(row["Ma_ban_ra"]).split(",")}
        trimmed = sorted(sold_symbols & target_symbols)
        removed = sorted(sold_symbols - target_symbols)
        sell_parts = []
        if removed:
            sell_parts.append(f"b\u00e1n to\u00e0n b\u1ed9 {', '.join(removed)}")
        if trimmed:
            sell_parts.append(f"b\u00e1n gi\u1ea3m t\u1ef7 tr\u1ecdng {', '.join(trimmed)}")
        sell_action = ", ".join(sell_parts) if sell_parts else f"b\u00e1n {row['Ma_ban_ra']}"
        action = (
            f"h\u1ec7 th\u1ed1ng quy\u1ebft \u0111\u1ecbnh c\u01a1 c\u1ea5u danh m\u1ee5c sang {row['Danh_muc_muc_tieu']}; "
            f"ng\u00e0y {execute_date}, {sell_action} v\u00e0 mua {row['Ma_mua_them']}, "
            f"gi\u00e1 tr\u1ecb mua {_with_accents(_money_vn(float(row['Gia_tri_mua'])))} v\u00e0 gi\u00e1 tr\u1ecb b\u00e1n "
            f"{_with_accents(_money_vn(float(row['Gia_tri_ban'])))}; chi ph\u00ed {fees}{realized}"
        )
    else:
        action = (
            f"h\u1ec7 th\u1ed1ng \u0111\u01b0a {row['Danh_muc_muc_tieu']} v\u00e0o danh m\u1ee5c m\u1ee5c ti\u00eau \u0111\u1ec3 m\u1edf mua; "
            f"ng\u00e0y {execute_date}, mua {row['Ma_mua_them']}, t\u1ed5ng gi\u00e1 tr\u1ecb mua "
            f"{_with_accents(_money_vn(float(row['Gia_tri_mua'])))}; chi ph\u00ed {fees}"
        )
    return (
        f"Ng\u00e0y {signal_date}, sau khi k\u1ebft th\u00fac phi\u00ean, {action}. "
        f"\u0110\u1ebfn k\u1ef3 \u0111\u00e1nh gi\u00e1 ng\u00e0y {end_date}, NAV c\u1ee7a chu k\u1ef3 {pnl} ({percent}); "
        f"danh m\u1ee5c cu\u1ed1i k\u1ef3 l\u00e0 {row['Nam_giu_cuoi_ky']}."
    )


def export(out_dir: Path) -> None:
    journal = pd.read_csv(out_dir / "flow_v2_nhat_ky_rebalance_von_1_ty.csv", encoding="utf-8-sig")
    annual = pd.read_csv(out_dir / "annual_breakdown.csv", encoding="utf-8-sig")
    realized = pd.read_csv(out_dir / "realized_roundtrips_fifo.csv", encoding="utf-8-sig")
    summary = pd.read_csv(out_dir / "summary.csv", encoding="utf-8-sig").set_index("metric")["value"].to_dict()
    baseline_equity = pd.read_csv(out_dir / "equity_curve_execution_corrected.csv")
    baseline_equity["date"] = pd.to_datetime(baseline_equity["date"])
    baseline_2024 = baseline_equity[
        baseline_equity["date"].between(pd.Timestamp("2024-01-01"), pd.Timestamp("2024-12-31"))
    ]["equity"]
    experiment = pd.read_csv(
        out_dir.parent / "flow_v2_reversal_guard_experiment_vn100_2020_to_2026-05-22" / "summary.csv"
    ).set_index("variant")
    risk_experiment = pd.read_csv(
        out_dir.parent / "flow_v2_risk_management_experiment_vn100_2020_to_2026-05-22" / "summary.csv"
    ).set_index("variant")
    risk_recent = pd.read_csv(
        out_dir.parent / "flow_v2_risk_management_experiment_vn100_2025_to_2026-05-22" / "summary.csv"
    ).set_index("variant")
    audit_dir = out_dir.parent / "flow_v2_detailed_audit_vn100_2020_to_2026-05-22"
    audit_full = pd.read_csv(audit_dir / "full_period_summary.csv").set_index("variant")
    audit_periods = pd.read_csv(audit_dir / "period_robustness_summary.csv")
    audit_checks = pd.read_csv(audit_dir / "execution_audit_checks.csv")
    trades = journal[(journal["So_lenh_mua"] + journal["So_lenh_ban"]) > 0].copy()
    quiet = journal[(journal["So_lenh_mua"] + journal["So_lenh_ban"]) == 0].copy()
    realized_by_date = realized.groupby("exit_date")["net_pnl"].sum().to_dict()

    event_text = [_describe(row, realized_by_date) for _, row in trades.iterrows()]
    annual_text = [
        (
            f"- N\u0103m {int(row['year'])}: t\u00e0i s\u1ea3n t\u0103ng/gi\u1ea3m "
            f"{_with_accents(_result_vn(float(row['pnl'])))} ({float(row['return_pct']):+.2f}%), "
            f"k\u1ebft th\u00fac n\u0103m \u1edf {_with_accents(_money_vn(float(row['ending_equity'])))}."
        )
        for _, row in annual.iterrows()
    ]
    text = "\n".join(
        [
            "B\u00c1O C\u00c1O DI\u1ec4N GI\u1ea2I L\u1ecaCH S\u1eec REBALANCE - MVP FLOW V2 HI\u1ec6U SU\u1ea4T CAO",
            "",
            "C\u1ea5u h\u00ecnh: VN100, high_rs + flow_heavy, t\u1ed1i \u0111a 2 m\u00e3, rebalance m\u1ed7i 10 phi\u00ean, market gate risk_on_or_strong_neutral.",
            "V\u1ed1n gi\u1ea3 l\u1eadp: 1 t\u1ef7 \u0111\u1ed3ng. Giai \u0111o\u1ea1n: 02/01/2020 \u0111\u1ebfn 22/05/2026.",
            "B\u1ea3n n\u00e0y gi\u1eef nguy\u00ean target l\u1ecbch s\u1eed \u0111\u00e3 kh\u00f3a v\u00e0 t\u00ednh l\u1ea1i th\u1ef1c thi v\u1edbi gi\u00e1 VND th\u1ef1c (gi\u00e1 parquet nh\u00e2n 1.000), l\u00f4 100 c\u1ed5 phi\u1ebfu.",
            "Chi ph\u00ed: m\u1ed9t deal mua r\u1ed3i b\u00e1n c\u00f3 chi ph\u00ed all-in x\u1ea5p x\u1ec9 0,40%: ph\u00eda mua g\u1ed3m 0,10% ph\u00ed v\u00e0 0,05% tr\u01b0\u1ee3t gi\u00e1; ph\u00eda b\u00e1n g\u1ed3m 0,10% ph\u00ed, 0,10% thu\u1ebf v\u00e0 0,05% tr\u01b0\u1ee3t gi\u00e1.",
            "Ng\u00e0y t\u00edn hi\u1ec7u l\u00e0 l\u00fac h\u1ec7 th\u1ed1ng ch\u1ed1t danh m\u1ee5c m\u1ee5c ti\u00eau sau \u0111\u00f3ng c\u1eeda; mua/b\u00e1n \u0111\u01b0\u1ee3c m\u00f4 ph\u1ecfng t\u1ea1i gi\u00e1 m\u1edf c\u1eeda phi\u00ean k\u1ebf ti\u1ebfp.",
            "",
            "LOGIC HO\u1ea0T \u0110\u1ed8NG C\u1ee6A MVP FLOW V2 HI\u1ec6U SU\u1ea4T CAO",
            "1. Pipeline d\u1eef li\u1ec7u: h\u1ec7 th\u1ed1ng gh\u00e9p gi\u00e1 OHLCV h\u00e0ng ng\u00e0y v\u1edbi Money Cycle (CHDM/DS), Smart Money Trace, s\u1ee9c m\u1ea1nh t\u01b0\u01a1ng \u0111\u1ed1i v\u00e0 thanh kho\u1ea3n \u0111\u1ec3 t\u1ea1o b\u1ea3ng feature cho t\u1eebng m\u00e3 trong VN100.",
            "2. Gate th\u1ecb tr\u01b0\u1eddng: h\u1ec7 th\u1ed1ng ch\u1ec9 l\u1eadp danh m\u1ee5c c\u1ed5 phi\u1ebfu khi tr\u1ea1ng th\u00e1i th\u1ecb tr\u01b0\u1eddng l\u00e0 RISK_ON ho\u1eb7c NEUTRAL, \u0111i\u1ec3m regime t\u1eeb 55 tr\u1edf l\u00ean, CHDM20 l\u1edbn h\u01a1n 48 v\u00e0 DS20 kh\u00f4ng qu\u00e1 0,50. Khi gate \u0111\u00f3ng, target l\u00e0 ti\u1ec1n m\u1eb7t.",
            "3. B\u1ed9 l\u1ecdc high_rs: m\u1ed7i m\u00e3 ph\u1ea3i n\u1eb1m tr\u00ean MA50, \u00e1p l\u1ef1c ph\u00e2n ph\u1ed1i kh\u00f4ng qu\u00e1 65, \u0111i\u1ec3m d\u00f2ng ti\u1ec1n t\u00e0i tr\u1ee3 t\u1eeb 45, gi\u00e1 tr\u1ecb giao d\u1ecbch so v\u1edbi trung b\u00ecnh 20 phi\u00ean t\u1eeb 0,80 v\u00e0 s\u1ee9c m\u1ea1nh t\u01b0\u01a1ng \u0111\u1ed1i 20 phi\u00ean n\u1eb1m trong nh\u00f3m tr\u00ean 60%.",
            "4. C\u00e1ch x\u1ebfp h\u1ea1ng flow_heavy: \u0111i\u1ec3m = 20% chu k\u1ef3 ng\u00e0nh + 30% d\u00f2ng ti\u1ec1n t\u00e0i tr\u1ee3 + 22% h\u1ea5p th\u1ee5 cung + 13% s\u1ee9c m\u1ea1nh t\u01b0\u01a1ng \u0111\u1ed1i + 10% x\u1ebfp h\u1ea1ng thanh kho\u1ea3n - 17% x\u1ebfp h\u1ea1ng \u00e1p l\u1ef1c ph\u00e2n ph\u1ed1i. Hai m\u00e3 c\u00f3 \u0111i\u1ec3m cao nh\u1ea5t sau khi qua gate s\u1ebd l\u00e0 target.",
            "5. Logic mua: c\u1ee9 m\u1ed7i 10 phi\u00ean giao d\u1ecbch, sau khi \u0111\u00f3ng c\u1eeda ng\u00e0y T, h\u1ec7 th\u1ed1ng ch\u1ed1t target; l\u1ec7nh mua/b\u00e1n \u0111\u01b0\u1ee3c m\u00f4 ph\u1ecfng t\u1ea1i gi\u00e1 m\u1edf c\u1eeda ng\u00e0y T+1. N\u1ebfu c\u00f3 2 target, danh m\u1ee5c \u0111\u01b0\u1ee3c c\u00e2n v\u1ec1 x\u1ea5p x\u1ec9 50% NAV cho m\u1ed7i m\u00e3, sau khi l\u00e0m tr\u00f2n theo l\u00f4 100 c\u1ed5 phi\u1ebfu.",
            "6. Logic b\u00e1n/rebalance: m\u00e3 kh\u00f4ng c\u00f2n trong target b\u1ecb b\u00e1n to\u00e0n b\u1ed9; m\u00e3 v\u1eabn thu\u1ed9c target c\u00f3 th\u1ec3 \u0111\u01b0\u1ee3c b\u00e1n b\u1edbt ho\u1eb7c mua th\u00eam \u0111\u1ec3 c\u00e2n l\u1ea1i t\u1ef7 tr\u1ecdng. N\u1ebfu gate \u0111\u00f3ng, h\u1ec7 th\u1ed1ng b\u00e1n danh m\u1ee5c v\u00e0 gi\u1eef ti\u1ec1n m\u1eb7t.",
            "7. Gi\u1edbi h\u1ea1n r\u1ee7i ro c\u1ee7a backtest n\u00e0y: b\u1ea3n hi\u1ec7u su\u1ea5t cao kh\u00f4ng c\u00f3 stop-loss ho\u1eb7c tho\u00e1t kh\u1ea9n c\u1ea5p gi\u1eefa hai k\u1ef3 rebalance. V\u00ec ch\u1ec9 n\u1eafm t\u1ed1i \u0111a 2 m\u00e3, m\u1ed9t l\u1ea7n xoay v\u00f2ng sai ng\u00e0nh c\u00f3 th\u1ec3 l\u00e0m NAV gi\u1ea3m m\u1ea1nh.",
            "8. Ph\u1ea1m vi c\u1ee7a b\u00e1o c\u00e1o: c\u00e1c quy\u1ebft \u0111\u1ecbnh target l\u1ecbch s\u1eed \u0111\u01b0\u1ee3c gi\u1eef nguy\u00ean; ph\u1ea7n kh\u1edbp l\u1ec7nh, s\u1ed1 c\u1ed5 phi\u1ebfu, ph\u00ed v\u00e0 NAV \u0111\u00e3 \u0111\u01b0\u1ee3c t\u00ednh l\u1ea1i b\u1eb1ng gi\u00e1 VND th\u1ef1c. \u0110\u00e2y kh\u00f4ng ph\u1ea3i l\u00e0 backtest c\u1ee7a l\u1edbp stop-loss/take-profit tr\u00ean dashboard paper.",
            "",
            "TH\u1eec NGHI\u1ec6M L\u1edaP CH\u1ed0NG \u0110\u1ea2O CHI\u1ec0U NG\u00c0NH",
            f"M\u1ed9t backtest t\u00edn hi\u1ec7u m\u1edbi \u0111\u00e3 th\u1eed gi\u1edbi h\u1ea1n m\u1ee9c t\u0103ng tr\u01b0\u1edbc khi mua, tho\u00e1t s\u1edbm khi momentum/flow g\u00e3y v\u00e0 so s\u00e1nh top2/top3/top4. Sau khi s\u1eeda l\u1ed7i gi\u1edbi h\u1ea1n v\u1ecb th\u1ebf, bi\u1ebfn th\u1ec3 tho\u00e1t s\u1edbm top2 \u0111\u1ea1t {float(experiment.loc['exit_top2', 'total_return_pct']):+.2f}%, Sharpe {float(experiment.loc['exit_top2', 'sharpe_ratio']):.3f} v\u00e0 max drawdown {float(experiment.loc['exit_top2', 'max_drawdown_pct']):.2f}%, so v\u1edbi baseline t\u00edn hi\u1ec7u m\u1edbi {float(experiment.loc['baseline_top2', 'total_return_pct']):+.2f}%, Sharpe {float(experiment.loc['baseline_top2', 'sharpe_ratio']):.3f} v\u00e0 max drawdown {float(experiment.loc['baseline_top2', 'max_drawdown_pct']):.2f}%.",
            "M\u1edf b\u00e1o c\u00e1o so s\u00e1nh: flow_v2_reversal_guard_experiment_vn100_2020_to_2026-05-22/index.html",
            "",
            "SO S\u00c1NH V\u1edaI FLOW V2 MVP BASELINE",
            "Flow V2 MVP baseline l\u00e0 k\u1ebft qu\u1ea3 report hi\u1ec7n t\u1ea1i v\u1edbi target l\u1ecbch s\u1eed \u0111\u00e3 kh\u00f3a; hai d\u00f2ng fresh-signal l\u00e0 backtest ch\u1ecdn l\u1ea1i danh m\u1ee5c t\u1eeb feature \u0111\u00e3 \u0111\u1ed3ng b\u1ed9, v\u00ec v\u1eady d\u00f9ng \u0111\u1ec3 so s\u00e1nh h\u00e0nh vi r\u1ee7i ro ch\u1ee9 kh\u00f4ng ph\u1ea3i thay th\u1ebf tr\u1ef1c ti\u1ebfp s\u1ed5 giao d\u1ecbch baseline.",
            "",
            "KI\u1ec2M TH\u1eec QU\u1ea2N TR\u1eca V\u1eca TH\u1ebe B\u1ed4 SUNG",
            f"K\u1ebft qu\u1ea3 t\u1ed1t nh\u1ea5t trong nh\u00f3m m\u1edbi l\u00e0 tho\u00e1t hai t\u1ea7ng top2: khi flow y\u1ebfu h\u1ec7 th\u1ed1ng gi\u1ea3m m\u1ed9t n\u1eeda v\u1ecb th\u1ebf, khi momentum/flow g\u00e3y m\u1ea1nh m\u1edbi b\u00e1n h\u1ebft \u1edf open T+1. Sau khi audit v\u00e0 s\u1eeda l\u1ed7i top2, c\u1ea5u h\u00ecnh n\u00e0y \u0111\u1ea1t {float(risk_experiment.loc['early_exit_tiered_top2', 'total_return_pct']):+.2f}%, Sharpe {float(risk_experiment.loc['early_exit_tiered_top2', 'sharpe_ratio']):.3f} v\u00e0 max drawdown {float(risk_experiment.loc['early_exit_tiered_top2', 'max_drawdown_pct']):.2f}%.",
            "M\u1edf b\u00e1o c\u00e1o qu\u1ea3n tr\u1ecb r\u1ee7i ro: flow_v2_risk_management_experiment_vn100_2020_to_2026-05-22/index.html",
            "",
            "BACKTEST KI\u1ec2M \u0110\u1ecaNH L\u1ea0I T\u1eea \u0110\u1ea6U",
            "L\u1ea7n audit m\u1edbi \u0111\u00e3 t\u00e1i ch\u1ea1y fresh-signal t\u1eeb d\u1eef li\u1ec7u ngu\u1ed3n, l\u01b0u ng\u00e0y t\u00edn hi\u1ec7u tr\u00ean t\u1eebng l\u1ec7nh v\u00e0 ki\u1ec3m tra kh\u1edbp open T+1. Audit ph\u00e1t hi\u1ec7n l\u1ed7i execution c\u0169: VIX kh\u00f4ng c\u00f3 gi\u00e1 m\u1edf c\u1eeda khi c\u1ea7n b\u00e1n ng\u00e0y 31/12/2020 nh\u01b0ng m\u00f4 ph\u1ecfng v\u1eabn m\u1edf th\u00eam CII v\u00e0 SSI, l\u00e0m danh m\u1ee5c t\u1ea1m th\u1eddi c\u00f3 3 m\u00e3. Engine \u0111\u00e3 \u0111\u01b0\u1ee3c s\u1eeda \u0111\u1ec3 v\u1ecb th\u1ebf ch\u01b0a b\u00e1n \u0111\u01b0\u1ee3c chi\u1ebfm m\u1ed9t slot, kh\u00f4ng mua v\u01b0\u1ee3t qu\u00e1 top2.",
            f"Sau s\u1eeda, baseline fresh-signal \u0111\u1ea1t {float(audit_full.loc['fresh_baseline_top2', 'total_return_pct']):+.2f}%, Sharpe {float(audit_full.loc['fresh_baseline_top2', 'sharpe_ratio']):.3f}, max drawdown {float(audit_full.loc['fresh_baseline_top2', 'max_drawdown_pct']):.2f}%; early-exit hai t\u1ea7ng top2 \u0111\u1ea1t {float(audit_full.loc['early_exit_tiered_top2', 'total_return_pct']):+.2f}%, Sharpe {float(audit_full.loc['early_exit_tiered_top2', 'sharpe_ratio']):.3f}, max drawdown {float(audit_full.loc['early_exit_tiered_top2', 'max_drawdown_pct']):.2f}%.",
            "K\u1ebft qu\u1ea3 kh\u00f4ng n\u00ean \u0111\u1ecdc m\u1ed9t chi\u1ec1u: khi kh\u1edfi t\u1ea1o l\u1ea1i v\u1ed1n ri\u00eang cho giai \u0111o\u1ea1n 2022-2023, baseline c\u00f3 l\u00e3i c\u00f2n early-exit hai t\u1ea7ng b\u1ecb l\u1ed7. Ngo\u00e0i ra, VN100 v\u1eabn l\u00e0 danh s\u00e1ch resolve t\u1ea1i th\u1eddi \u0111i\u1ec3m ch\u1ea1y, ch\u01b0a lo\u1ea1i b\u1ecf ho\u00e0n to\u00e0n survivorship bias.",
            "M\u1edf b\u00e1o c\u00e1o audit chi ti\u1ebft: flow_v2_detailed_audit_vn100_2020_to_2026-05-22/index.html",
            "",
            "T\u1ed4NG QUAN",
            f"- V\u1ed1n ban \u0111\u1ea7u 1 t\u1ef7 \u0111\u1ed3ng, t\u00e0i s\u1ea3n cu\u1ed1i k\u1ef3 {_with_accents(_money_vn(float(summary['Ending equity'])))}, l\u00e3i r\u00f2ng {_with_accents(_money_vn(float(summary['Net P&L'])))} ({float(summary['Total return (%)']):+.2f}%).",
            f"- C\u00f3 {int(float(summary['Executed rebalances']))} k\u1ef3 rebalance \u0111\u00e3 th\u1ef1c thi; {int(float(summary['Target-bearing windows']))} k\u1ef3 c\u00f3 danh m\u1ee5c c\u1ed5 phi\u1ebfu m\u1ee5c ti\u00eau.",
            f"- Trong c\u00e1c k\u1ef3 c\u00f3 target: {int(float(summary['Target windows won']))} k\u1ef3 l\u00e3i, {int(float(summary['Target windows lost']))} k\u1ef3 l\u1ed7. Max drawdown l\u00e0 {float(summary['Max drawdown (%)']):.2f}%.",
            "",
            "HI\u1ec6U QU\u1ea2 THEO N\u0102M",
            *annual_text,
            "",
            "DI\u1ec4N BI\u1ebeN C\u00c1C L\u1ea6N C\u00d3 GIAO D\u1ecaCH",
            *[f"\nL\u1ea7n rebalance {int(row['Lan_rebalance'])}: {event}" for (_, row), event in zip(trades.iterrows(), event_text)],
            "",
            "C\u00c1C K\u1ef2 KH\u00d4NG C\u00d3 L\u1ec6NH",
            f"- C\u00f3 {len(quiet)} k\u1ef3 h\u1ec7 th\u1ed1ng kh\u00f4ng mua/b\u00e1n, ch\u1ee7 y\u1ebfu do gate \u0111\u00f3ng v\u00e0 danh m\u1ee5c gi\u1eef ti\u1ec1n m\u1eb7t.",
            "",
            "L\u01afU \u00dd KI\u1ec2M \u0110\u1ecaNH",
            "- B\u00e1o c\u00e1o n\u00e0y l\u00e0 backtest v\u1ed1n 1 t\u1ef7; dashboard hi\u1ec7n t\u1ea1i v\u1eabn d\u00f9ng v\u1ed1n sleeve 100 tri\u1ec7u.",
            "- K\u1ebft qu\u1ea3 c\u00f2n r\u1ee7i ro survivorship bias do d\u00f9ng danh s\u00e1ch VN100 hi\u1ec7n t\u1ea1i cho l\u1ecbch s\u1eed.",
        ]
    )
    (out_dir / "bao_cao_dien_giai_rebalance_von_1_ty.txt").write_text(text, encoding="utf-8")
    (out_dir / "bao_cao_dien_giai_rebalance_von_1_ty.md").write_text(text, encoding="utf-8")

    headings = {
        "LOGIC HO\u1ea0T \u0110\u1ed8NG C\u1ee6A MVP FLOW V2 HI\u1ec6U SU\u1ea4T CAO",
        "TH\u1eec NGHI\u1ec6M L\u1edaP CH\u1ed0NG \u0110\u1ea2O CHI\u1ec0U NG\u00c0NH",
        "SO S\u00c1NH V\u1edaI FLOW V2 MVP BASELINE",
        "KI\u1ec2M TH\u1eec QU\u1ea2N TR\u1eca V\u1eca TH\u1ebe B\u1ed4 SUNG",
        "BACKTEST KI\u1ec2M \u0110\u1ecaNH L\u1ea0I T\u1eea \u0110\u1ea6U",
        "T\u1ed4NG QUAN",
        "HI\u1ec6U QU\u1ea2 THEO N\u0102M",
        "DI\u1ec4N BI\u1ebeN C\u00c1C L\u1ea6N C\u00d3 GIAO D\u1ecaCH",
        "C\u00c1C K\u1ef2 KH\u00d4NG C\u00d3 L\u1ec6NH",
        "L\u01afU \u00dd KI\u1ec2M \u0110\u1ecaNH",
    }
    fresh_baseline = experiment.loc["baseline_top2"]
    early_exit = experiment.loc["exit_top2"]
    comparison_rows = [
        (
            "Flow V2 MVP baseline",
            "Target l\u1ecbch s\u1eed \u0111\u00e3 kh\u00f3a, gi\u00e1 VND \u0111\u00e3 s\u1eeda",
            float(summary["Total return (%)"]),
            float(summary["Sharpe"]),
            float(summary["Max drawdown (%)"]),
            float(annual.loc[annual["year"] == 2024, "return_pct"].iloc[0]),
            _max_drawdown_pct(baseline_2024),
        ),
        (
            "Fresh signal baseline top2",
            "Ch\u1ecdn l\u1ea1i target, ch\u01b0a c\u00f3 early exit",
            float(fresh_baseline["total_return_pct"]),
            float(fresh_baseline["sharpe_ratio"]),
            float(fresh_baseline["max_drawdown_pct"]),
            float(fresh_baseline["return_2024_pct"]),
            float(fresh_baseline["drawdown_2024_pct"]),
        ),
        (
            "Fresh signal early-exit top2",
            "Ch\u1ecdn l\u1ea1i target, tho\u00e1t open T+1 khi flow/momentum g\u00e3y",
            float(early_exit["total_return_pct"]),
            float(early_exit["sharpe_ratio"]),
            float(early_exit["max_drawdown_pct"]),
            float(early_exit["return_2024_pct"]),
            float(early_exit["drawdown_2024_pct"]),
        ),
    ]
    comparison_html = (
        '<table class="comparison"><thead><tr><th>C\u1ea5u h\u00ecnh</th><th>L\u1ee3i nhu\u1eadn to\u00e0n k\u1ef3</th>'
        "<th>Sharpe</th><th>Max DD</th><th>L\u1ee3i nhu\u1eadn 2024</th><th>DD 2024</th></tr></thead><tbody>"
        + "".join(
            "<tr>"
            f"<td><strong>{escape(name)}</strong><br><span>{escape(detail)}</span></td>"
            f"<td>{return_pct:+.2f}%</td><td>{sharpe:.3f}</td><td>{drawdown:.2f}%</td>"
            f"<td>{return_2024:+.2f}%</td><td>{drawdown_2024:.2f}%</td>"
            "</tr>"
            for name, detail, return_pct, sharpe, drawdown, return_2024, drawdown_2024 in comparison_rows
        )
        + "</tbody></table>"
    )
    risk_labels = {
        "fresh_baseline_top2": ("Fresh signal baseline top2", "Kh\u00f4ng overlay qu\u1ea3n tr\u1ecb v\u1ecb th\u1ebf"),
        "early_exit_full_top2": ("Early-exit full top2", "G\u00e3y flow/momentum th\u00ec b\u00e1n h\u1ebft"),
        "early_exit_tiered_top2": ("Early-exit hai t\u1ea7ng top2", "Y\u1ebfu th\u00ec gi\u1ea3m n\u1eeda; g\u00e3y m\u1ea1nh th\u00ec b\u00e1n h\u1ebft"),
        "early_exit_inverse_atr_top2": ("Early-exit + inverse ATR", "T\u1ef7 tr\u1ecdng th\u1ea5p h\u01a1n cho m\u00e3 bi\u1ebfn \u0111\u1ed9ng cao"),
        "early_exit_sector_cap_top2": ("Early-exit + kh\u00e1c ng\u00e0nh", "T\u1ed1i \u0111a m\u1ed9t m\u00e3 m\u1ed7i ng\u00e0nh"),
        "early_exit_sideway_buffer_top2": ("Early-exit + NEUTRAL 70%", "Gate NEUTRAL ch\u1ec9 gi\u1ea3i ng\u00e2n 70% NAV"),
        "combined_risk_overlay_top2": ("K\u1ebft h\u1ee3p t\u1ea5t c\u1ea3 overlay", "Hai t\u1ea7ng + ATR + kh\u00e1c ng\u00e0nh + NEUTRAL 70%"),
    }
    risk_rows = [
        (
            "Flow V2 MVP baseline",
            "Target l\u1ecbch s\u1eed \u0111\u00e3 kh\u00f3a",
            float(summary["Total return (%)"]),
            float(summary["Sharpe"]),
            float(summary["Max drawdown (%)"]),
            float(annual.loc[annual["year"] == 2024, "return_pct"].iloc[0]),
            "-",
        )
    ]
    for variant in risk_labels:
        name, detail = risk_labels[variant]
        row = risk_experiment.loc[variant]
        recent = risk_recent.loc[variant]
        risk_rows.append(
            (
                name,
                detail,
                float(row["total_return_pct"]),
                float(row["sharpe_ratio"]),
                float(row["max_drawdown_pct"]),
                float(row["return_2024_pct"]),
                f"{float(recent['total_return_pct']):+.2f}%",
            )
        )
    risk_table_html = (
        '<table class="comparison"><thead><tr><th>C\u1ea5u h\u00ecnh</th><th>L\u1ee3i nhu\u1eadn to\u00e0n k\u1ef3</th>'
        "<th>Sharpe</th><th>Max DD</th><th>N\u0103m 2024</th><th>Test t\u1eeb 01/01/2025</th></tr></thead><tbody>"
        + "".join(
            "<tr>"
            f"<td><strong>{escape(name)}</strong><br><span>{escape(detail)}</span></td>"
            f"<td>{return_pct:+.2f}%</td><td>{sharpe:.3f}</td><td>{drawdown:.2f}%</td>"
            f"<td>{return_2024:+.2f}%</td><td>{escape(recent_return)}</td>"
            "</tr>"
            for name, detail, return_pct, sharpe, drawdown, return_2024, recent_return in risk_rows
        )
        + "</tbody></table>"
    )
    audit_period_labels = {
        "2020-2021 khoi phuc/tang manh": "2020-2021: ph\u1ee5c h\u1ed3i / t\u0103ng m\u1ea1nh",
        "2022-2023 dieu chinh/hoi phuc": "2022-2023: \u0111i\u1ec1u ch\u1ec9nh / h\u1ed3i ph\u1ee5c",
        "2024 sideway stress": "2024: ki\u1ec3m tra sideway",
        "2025-den-22/05/2026 gan day": "2025 \u0111\u1ebfn 22/05/2026: g\u1ea7n \u0111\u00e2y",
        "2020-2021 ph\u1ee5c h\u1ed3i/t\u0103ng m\u1ea1nh": "2020-2021: ph\u1ee5c h\u1ed3i / t\u0103ng m\u1ea1nh",
        "2022-2023 \u0111i\u1ec1u ch\u1ec9nh/h\u1ed3i ph\u1ee5c": "2022-2023: \u0111i\u1ec1u ch\u1ec9nh / h\u1ed3i ph\u1ee5c",
        "2024 ki\u1ec3m tra sideway": "2024: ki\u1ec3m tra sideway",
        "2025 \u0111\u1ebfn 22/05/2026 g\u1ea7n \u0111\u00e2y": "2025 \u0111\u1ebfn 22/05/2026: g\u1ea7n \u0111\u00e2y",
    }
    audit_rows = []
    for variant in ["fresh_baseline_top2", "early_exit_tiered_top2"]:
        row = audit_full.loc[variant]
        audit_name = (
            "Baseline fresh-signal top2"
            if variant == "fresh_baseline_top2"
            else "Early-exit hai t\u1ea7ng top2"
        )
        audit_rows.append(
            "<tr>"
            f"<td><strong>{audit_name}</strong></td>"
            f"<td>{float(row['total_return_pct']):+.2f}%</td><td>{float(row['sharpe_ratio']):.3f}</td>"
            f"<td>{float(row['max_drawdown_pct']):.2f}%</td><td>{int(row['number_of_orders'])}</td>"
            f"<td>{_with_accents(_money_vn(float(row['total_fees'])))}</td></tr>"
        )
    audit_table_html = (
        '<table class="comparison"><thead><tr><th>C\u1ea5u h\u00ecnh audit sau s\u1eeda</th><th>L\u1ee3i nhu\u1eadn</th>'
        "<th>Sharpe</th><th>Max DD</th><th>S\u1ed1 l\u1ec7nh</th><th>Ph\u00ed/thu\u1ebf</th></tr></thead><tbody>"
        + "".join(audit_rows)
        + "</tbody></table>"
    )
    audit_period_html = (
        '<table class="comparison"><thead><tr><th>Giai \u0111o\u1ea1n kh\u1edfi t\u1ea1o l\u1ea1i v\u1ed1n</th><th>C\u1ea5u h\u00ecnh</th>'
        "<th>L\u1ee3i nhu\u1eadn</th><th>Sharpe</th><th>Max DD</th></tr></thead><tbody>"
        + "".join(
            "<tr>"
            f"<td>{audit_period_labels.get(str(row['period']), escape(str(row['period'])))}</td>"
            f"<td>{'Baseline' if row['variant'] == 'fresh_baseline_top2' else 'Early-exit hai tầng'}</td>"
            f"<td>{float(row['total_return_pct']):+.2f}%</td><td>{float(row['sharpe_ratio']):.3f}</td>"
            f"<td>{float(row['max_drawdown_pct']):.2f}%</td></tr>"
            for _, row in audit_periods.iterrows()
        )
        + "</tbody></table>"
    )
    body_parts = []
    for line in text.splitlines():
        if not line:
            continue
        if line in headings:
            body_parts.append(f"<h2>{escape(line)}</h2>")
            if line == "SO S\u00c1NH V\u1edaI FLOW V2 MVP BASELINE":
                body_parts.append(comparison_html)
            elif line == "KI\u1ec2M TH\u1eec QU\u1ea2N TR\u1eca V\u1eca TH\u1ebe B\u1ed4 SUNG":
                body_parts.append(risk_table_html)
            elif line == "BACKTEST KI\u1ec2M \u0110\u1ecaNH L\u1ea0I T\u1eea \u0110\u1ea6U":
                body_parts.extend([audit_table_html, audit_period_html])
        elif line.startswith("M\u1edf b\u00e1o c\u00e1o so s\u00e1nh:"):
            body_parts.append(
                '<p><a href="../flow_v2_reversal_guard_experiment_vn100_2020_to_2026-05-22/index.html">'
                "M\u1edf b\u00e1o c\u00e1o so s\u00e1nh chi ti\u1ebft c\u00e1c overlay"
                "</a></p>"
            )
        elif line.startswith("M\u1edf b\u00e1o c\u00e1o qu\u1ea3n tr\u1ecb r\u1ee7i ro:"):
            body_parts.append(
                '<p><a href="../flow_v2_risk_management_experiment_vn100_2020_to_2026-05-22/index.html">'
                "M\u1edf b\u00e1o c\u00e1o chi ti\u1ebft overlay qu\u1ea3n tr\u1ecb v\u1ecb th\u1ebf"
                "</a></p>"
            )
        elif line.startswith("M\u1edf b\u00e1o c\u00e1o audit chi ti\u1ebft:"):
            body_parts.append(
                '<p><a href="../flow_v2_detailed_audit_vn100_2020_to_2026-05-22/index.html">'
                "M\u1edf b\u00e1o c\u00e1o backtest ki\u1ec3m \u0111\u1ecbnh l\u1ea1i t\u1eeb \u0111\u1ea7u"
                "</a></p>"
            )
        elif line.startswith("L\u1ea7n rebalance "):
            body_parts.append(f"<p class=\"event\">{escape(line)}</p>")
        else:
            body_parts.append(f"<p>{escape(line)}</p>")
    paragraphs = "".join(body_parts)
    html = (
        '<!doctype html><html lang="vi"><head><meta charset="utf-8">'
        "<title>B\u00e1o c\u00e1o di\u1ec5n gi\u1ea3i Flow V2 - v\u1ed1n 1 t\u1ef7</title>"
        "<style>body{font:15px Segoe UI,Arial,sans-serif;background:#f4f7f8;color:#16242b;margin:0}"
        "main{background:white;max-width:1040px;margin:24px auto;padding:28px;border:1px solid #dce4e6;border-radius:6px}"
        "p{line-height:1.65;margin:9px 0;white-space:pre-wrap}p:first-child{font-size:25px;font-weight:700;color:#146b62}"
        "h2{font-size:18px;letter-spacing:0;color:#146b62;margin:26px 0 12px;border-top:1px solid #e3ebec;padding-top:18px}"
        ".event{padding:9px 12px;background:#f6f9f9;border-left:3px solid #b5d4cf;margin:11px 0}"
        "table.comparison{width:100%;border-collapse:collapse;margin:12px 0 16px;font-size:14px}"
        "table.comparison th,table.comparison td{border:1px solid #dce4e6;padding:10px;text-align:right;vertical-align:top}"
        "table.comparison th{background:#edf4f3;color:#214342}table.comparison th:first-child,table.comparison td:first-child{text-align:left}"
        "table.comparison span{font-size:12px;color:#596c74}"
        "a{color:#146b62;font-weight:600;text-decoration:none}a:hover{text-decoration:underline}"
        "</style></head><body><main>" + paragraphs + "</main></body></html>"
    )
    (out_dir / "bao_cao_dien_giai_rebalance_von_1_ty.html").write_text(html, encoding="utf-8")

    source_book = out_dir / "flow_v2_nhat_ky_rebalance_von_1_ty.xlsx"
    dest_book = out_dir / "flow_v2_nhat_ky_rebalance_von_1_ty_co_dien_giai_phi_04deal.xlsx"
    workbook = load_workbook(source_book)
    if "Dien giai de doc" in workbook.sheetnames:
        del workbook["Dien giai de doc"]
    sheet = workbook.create_sheet("Dien giai de doc", 1)
    sheet.column_dimensions["A"].width = 145
    sheet["A1"] = "B\u00e1o c\u00e1o di\u1ec5n gi\u1ea3i Flow V2 - chi ph\u00ed all-in kho\u1ea3ng 0,40% / deal mua-b\u00e1n"
    sheet["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="146B62")
    sheet["A1"].alignment = Alignment(wrap_text=True)
    for index, line in enumerate(text.splitlines()[2:], start=3):
        cell = sheet.cell(row=index, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[index].height = 60 if line.startswith("L\u1ea7n rebalance") else 26
    sheet.freeze_panes = "A3"
    workbook.save(dest_book)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    export(args.out_dir)


if __name__ == "__main__":
    main()
